"""
Azure Function App for Document and Excel Processing

This application provides multiple endpoints for processing text documents and Excel files
using LangGraph pipelines with Azure OpenAI integration.

Available Endpoints:
    - POST /process-text: Synchronous text processing
    - POST /process-excel: Synchronous Excel file processing
    - POST /generate-upload-url: Generate SAS URLs for file uploads
    - Blob trigger: Automatic processing of uploaded files

Architecture:
    - LangGraph state machine for agent orchestration
    - Azure OpenAI Chat Completions API for AI processing
    - Azure Blob Storage for file management
    - Azure Durable Functions for asynchronous workflows
"""

import azure.functions as func
import azure.durable_functions as df
import json
import logging
import os
import traceback
import importlib.util
from io import BytesIO
from datetime import datetime, timedelta, timezone
from typing import Dict, Any

from azure.storage.blob import (
    BlobServiceClient,
    generate_blob_sas,
    BlobSasPermissions,
    ContainerClient,
    ContentSettings,
)

# Import LangGraph pipelines from reorganized agent directories
from agents.text.langgraph_pipeline import run_langgraph_document_pipeline
from agents.text.parse_text_agent import AgentState
from agents.excel.excel_langgraph_pipeline import run_excel_langgraph_pipeline

# Configure logging
logging.basicConfig(level=logging.INFO)

# Initialize Function App
app = func.FunctionApp(http_auth_level=func.AuthLevel.ANONYMOUS)

# Constants
WORKFLOW_VERSION = "3.2"
ALLOWED_TEXT_EXTENSIONS = (".txt", ".docx")
ALLOWED_EXCEL_EXTENSIONS = (".xlsx", ".xls")
ALLOWED_ALL_EXTENSIONS = ALLOWED_TEXT_EXTENSIONS + ALLOWED_EXCEL_EXTENSIONS
SAS_URL_EXPIRY_MINUTES = 15
STORAGE_CONTAINER_UPLOADS = "uploads"
STORAGE_CONTAINER_RESULTS = "results"



# ==============================================================================
# HELPER FUNCTIONS
# ==============================================================================

def _load_pdf_converter():
    """
    Dynamically load the PDF converter utility module.
    
    Returns:
        module: The loaded pdf_converter module
        
    Raises:
        ImportError: If the module cannot be loaded
    """
    repo_root = os.path.dirname(os.path.dirname(__file__))
    converter_path = os.path.join(repo_root, "utilities-functions", "pdf_converter.py")

    spec = importlib.util.spec_from_file_location("pdf_converter", converter_path)
    if spec is None or spec.loader is None:
        raise ImportError(f"Unable to load pdf converter from {converter_path}")

    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _extract_text(blob_bytes: bytes, filename: str) -> str:
    """
    Extract plain text content from .txt or .docx file bytes.
    
    Args:
        blob_bytes: Raw file content as bytes
        filename: Name of the file (used to determine type)
        
    Returns:
        str: Extracted text content
        
    Raises:
        ValueError: If file type is not supported
    """
    if filename.endswith(".txt"):
        return blob_bytes.decode("utf-8")
    
    elif filename.endswith(".docx"):
        import io
        from docx import Document
        
        logging.info(f"Parsing .docx content for: {filename}")
        doc = Document(io.BytesIO(blob_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
        logging.info(f"Extracted {len(paragraphs)} non-empty paragraphs from {filename}")
        
        return "\n".join(paragraphs)
    
    else:
        raise ValueError(f"Unsupported file type: {filename}")


def _run_agent_pipeline(text: str, filename: str = "document.docx") -> dict:
    """
    Execute the LangGraph document processing pipeline.
    
    Args:
        text: Document text to process
        filename: Name of the document file
        
    Returns:
        dict: Final state from the LangGraph pipeline containing processed results
    """
    return run_langgraph_document_pipeline(text, filename)


def _build_docx_bytes_from_text(text: str) -> bytes:
    """
    Create a .docx document from plain text, preserving paragraph structure.
    
    Args:
        text: Text content to convert to .docx format
        
    Returns:
        bytes: Binary content of the created .docx file
    """
    from docx import Document

    document = Document()
    
    lines = text.split("\n")
    if not lines:
        lines = [text]

    for line in lines:
        document.add_paragraph(line)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ==============================================================================
# HTTP ENDPOINTS - Synchronous Processing
# ==============================================================================
@app.route(route="process-text", methods=["POST"])
def http_start(req: func.HttpRequest) -> func.HttpResponse:
    """
    Synchronous HTTP endpoint for text document processing.
    
    Accepts JSON with 'text' field and optional 'filename' field.
    Processes the text through the LangGraph pipeline with Azure OpenAI.
    
    Request Body:
        {
            "text": "document content here",
            "filename": "optional_name.txt"
        }
        
    Returns:
        JSON response with processing results including title, summary, 
        synonyms, enhanced text, and token usage statistics
    """
    try:
        req_body = req.get_json()
        if not req_body or "text" not in req_body:
            return func.HttpResponse(
                json.dumps({"error": "Missing 'text' in request body"}),
                status_code=400,
                mimetype="application/json",
            )

        text = req_body.get("text")
        filename = req_body.get("filename", "document.txt")
        
        logging.info(f"Processing text from /process-text: {text[:100]}...")

        # Execute LangGraph pipeline for text processing
        final_state = _run_agent_pipeline(text, filename)
        
        logging.info(f"Processing completed - Title: {final_state.get('title', 'N/A')}")
        logging.info(f"Token usage: {final_state.get('token_usage', {})}")

        response = {
            "success": True,
            "original_text": final_state.get("text", ""),
            "enhanced_text": final_state.get("enhanced_text", ""),
            "title": final_state.get("title", ""),
            "summary": final_state.get("summary", ""),
            "synonyms_found": len(final_state.get("synonyms", {})),
            "synonyms_applied": final_state.get("synonyms", {}),
            "token_usage": final_state.get("token_usage", {}),
            "pdf_base64": final_state.get("pdf_content", ""),
            "parsed_data": final_state.get("parsed_data", {}),
            "workflow_info": {
                "agents_used": [
                    "parse_text",
                    "find_equivalents",
                    "generate_summary",
                    "consolidate_text"
                ],
                "workflow_type": "langgraph_chat_completions_api_pipeline",
                "version": WORKFLOW_VERSION,
                "langgraph_enabled": True,
                "azure_openai_chat_completions_enabled": True,
            },
        }

        return func.HttpResponse(
            json.dumps(response, ensure_ascii=False),
            status_code=200,
            mimetype="application/json",
        )

    except Exception as e:
        logging.error(f"HTTP trigger error: {str(e)}")
        return func.HttpResponse(
            json.dumps({"error": str(e)}),
            status_code=500,
            mimetype="application/json",
        )


@app.route(route="process-excel", methods=["POST"])
def process_excel_http(req: func.HttpRequest) -> func.HttpResponse:
    """
    Synchronous HTTP endpoint for Excel file processing.
    
    Accepts either multipart form data with file upload or JSON with base64-encoded content.
    Processes Excel files through the LangGraph pipeline with analysis and RAGAS evaluation.
    
    Request Options:
        1. Multipart form data: file field with .xlsx or .xls file
        2. JSON body: {"excel_content": "base64...", "filename": "file.xlsx"}
        
    Returns:
        JSON response with parsing results, analysis, title, summary, and RAGAS scores
    """
    try:
        # Support multiple input formats for flexibility
        content_type = req.headers.get('Content-Type', '')
        
        if 'multipart/form-data' in content_type:
            # Handle file upload
            files = req.files
            if 'file' not in files:
                return func.HttpResponse(
                    json.dumps({"error": "No file provided in multipart request"}),
                    status_code=400,
                    mimetype="application/json",
                )
            
            file = files['file']
            excel_content = file.read()
            filename = file.filename or "uploaded.xlsx"
            
        else:
            # Handle JSON with base64 content
            req_body = req.get_json()
            if not req_body or "excel_content" not in req_body:
                return func.HttpResponse(
                    json.dumps({"error": "Missing 'excel_content' in request body"}),
                    status_code=400,
                    mimetype="application/json",
                )

            import base64
            excel_content = base64.b64decode(req_body.get("excel_content"))
            filename = req_body.get("filename", "uploaded.xlsx")
        
        # Ensure only Excel files are processed
        if not filename.lower().endswith(ALLOWED_EXCEL_EXTENSIONS):
            return func.HttpResponse(
                json.dumps({"error": "Only .xlsx and .xls files are supported"}),
                status_code=400,
                mimetype="application/json",
            )
        
        logging.info(
            f"Processing Excel file: {filename} "
            f"({len(excel_content)} bytes)"
        )

        # Execute Excel LangGraph pipeline
        result = run_excel_langgraph_pipeline(excel_content, filename)
        
        logging.info(f"Excel processing completed: {filename}")
        logging.info(f"Generated title: {result.get('title', 'N/A')}")
        logging.info(f"Token usage: {result.get('token_usage', {})}")

        return func.HttpResponse(
            json.dumps(result, ensure_ascii=False),
            status_code=200,
            mimetype="application/json",
        )

    except Exception as e:
        logging.error(f"Excel processing HTTP trigger error: {str(e)}")
        return func.HttpResponse(
            json.dumps({"error": str(e)}),
            status_code=500,
            mimetype="application/json",
        )


# ==============================================================================
# DURABLE FUNCTIONS - Asynchronous Orchestration
# ==============================================================================
@app.orchestration_trigger(context_name="context")
def text_processing_orchestrator(context: df.DurableOrchestrationContext):
    """
    Durable Functions orchestrator for asynchronous text processing.
    
    Coordinates the execution of the LangGraph pipeline as a single activity.
    Provides better reliability and monitoring for long-running operations.
    
    Args:
        context: Durable orchestration context with input data
        
    Returns:
        dict: Processing results with workflow metadata
    """
    try:
        input_data = context.get_input()
        text = input_data.get("text", "")
        filename = input_data.get("filename", "document.txt")

        logging.info(f"LangGraph orchestrator started for: {filename}")

        initial_state = {
            "text": text,
            "filename": filename
        }

        # Single activity call that runs the entire LangGraph
        final_state = yield context.call_activity("langgraph_pipeline_activity", initial_state)

        response = {
            "success": True,
            "original_text": final_state.get("text", ""),
            "enhanced_text": final_state.get("enhanced_text", ""),
            "title": final_state.get("title", ""),
            "summary": final_state.get("summary", ""),
            "synonyms_found": len(final_state.get("synonyms", {})),
            "synonyms_applied": final_state.get("synonyms", {}),
            "token_usage": final_state.get("token_usage", {}),
            "pdf_base64": final_state.get("pdf_content", ""),
            "parsed_data": final_state.get("parsed_data", {}),
            "workflow_info": {
                "agents_used": [
                    "parse_text",
                    "find_equivalents",
                    "generate_summary",
                    "consolidate_text"
                ],
                "workflow_type": "langgraph_chat_completions_durable_functions",
                "version": WORKFLOW_VERSION,
                "instance_id": context.instance_id,
                "langgraph_enabled": True,
                "azure_openai_chat_completions_enabled": True,
            },
        }

        logging.info(f"LangGraph orchestration completed for instance: {context.instance_id}")
        return response

    except Exception as e:
        logging.error(f"LangGraph orchestrator error: {str(e)}")
        return {"error": str(e), "workflow": "langgraph_durable_functions"}


@app.activity_trigger(input_name="input")
def langgraph_pipeline_activity(input: dict) -> dict:
    """
    Durable Functions activity that executes the LangGraph pipeline.
    
    This activity function wraps the LangGraph execution for use in
    durable orchestrations, providing retry and monitoring capabilities.
    
    Args:
        input: Dictionary with 'text' and 'filename' fields
        
    Returns:
        dict: Complete pipeline results
        
    Raises:
        Exception: Any errors during pipeline execution
    """
    logging.info("Executing LangGraph pipeline activity")
    try:
        text = input.get("text", "")
        filename = input.get("filename", "document.txt")
        result = run_langgraph_document_pipeline(text, filename)
        logging.info("LangGraph pipeline activity completed successfully")
        return result
    except Exception as e:
        logging.error(f"LangGraph pipeline activity error: {str(e)}")
        raise



# ==============================================================================
# LEGACY ACTIVITY FUNCTIONS - Backward Compatibility
# ==============================================================================

@app.activity_trigger(input_name="input")
def parse_text_activity(input: dict) -> dict:
    """
    Legacy activity for text parsing (deprecated).
    
    Note: This function is maintained for backward compatibility only.
    New implementations should use langgraph_pipeline_activity instead.
    """
    logging.info(
        "Executing legacy parse_text_activity - "
        "consider using langgraph_pipeline_activity"
    )
    try:
        from agents.text.parse_text_agent import parse_text_agent
        result = parse_text_agent(input)
        logging.info("parse_text_activity completed successfully")
        return result
    except Exception as e:
        logging.error(f"parse_text_activity error: {str(e)}")
        raise


@app.activity_trigger(input_name="input")
def find_equivalents_activity(input: dict) -> dict:
    """
    Legacy activity for synonym finding (deprecated).
    
    Note: This function is maintained for backward compatibility only.
    New implementations should use langgraph_pipeline_activity instead.
    """
    logging.info(
        "Executing legacy find_equivalents_activity - "
        "consider using langgraph_pipeline_activity"
    )
    try:
        from agents.text.find_equivalents_agent import find_equivalents_agent
        result = find_equivalents_agent(input)
        logging.info("find_equivalents_activity completed successfully")
        return result
    except Exception as e:
        logging.error(f"find_equivalents_activity error: {str(e)}")
        raise


@app.activity_trigger(input_name="input")
def consolidate_text_activity(input: dict) -> dict:
    """
    Legacy activity for text consolidation (deprecated).
    
    Note: This function is maintained for backward compatibility only.
    New implementations should use langgraph_pipeline_activity instead.
    """
    logging.info(
        "Executing legacy consolidate_text_activity - "
        "consider using langgraph_pipeline_activity"
    )
    try:
        from agents.text.consolidate_text_agent import consolidate_text_agent
        result = consolidate_text_agent(input)
        logging.info("consolidate_text_activity completed successfully")
        return result
    except Exception as e:
        logging.error(f"consolidate_text_activity error: {str(e)}")
        raise


# ==============================================================================
# FILE UPLOAD SUPPORT - SAS URL Generation
# ==============================================================================
@app.route(route="generate-upload-url", methods=["POST"])
def generate_upload_url(req: func.HttpRequest) -> func.HttpResponse:
    """
    Generate SAS URL for secure client-side file uploads.
    
    Creates a time-limited, write-only SAS URL that allows clients to upload
    files directly to Azure Blob Storage without requiring credentials.
    
    Request Body:
        {
            "filename": "document.xlsx"
        }
        
    Supported Extensions:
        .txt, .docx, .xlsx, .xls
        
    Returns:
        JSON with SAS URL, blob name, container, and expiry time
    """
    try:
        req_body = req.get_json()
        if not req_body or "filename" not in req_body:
            return func.HttpResponse(
                json.dumps({"error": "Missing 'filename' in request body"}),
                status_code=400,
                mimetype="application/json",
            )

        filename = req_body["filename"]
        
        if not filename.endswith(ALLOWED_ALL_EXTENSIONS):
            return func.HttpResponse(
                json.dumps({
                    "error": f"Only {', '.join(ALLOWED_ALL_EXTENSIONS)} files are supported"
                }),
                status_code=400,
                mimetype="application/json",
            )

        account_name = os.environ["STORAGE_ACCOUNT_NAME"]
        account_key = os.environ["STORAGE_ACCOUNT_KEY"]
        container_name = STORAGE_CONTAINER_UPLOADS

        expiry = datetime.now(timezone.utc) + timedelta(minutes=SAS_URL_EXPIRY_MINUTES)

        sas_token = generate_blob_sas(
            account_name=account_name,
            container_name=container_name,
            blob_name=filename,
            account_key=account_key,
            permission=BlobSasPermissions(write=True, create=True),
            expiry=expiry,
        )

        sas_url = (
            f"https://{account_name}.blob.core.windows.net"
            f"/{container_name}/{filename}?{sas_token}"
        )

        logging.info(f"Generated SAS URL for blob: {filename}")

        return func.HttpResponse(
            json.dumps({
                "sas_url": sas_url,
                "blob_name": filename,
                "container": container_name,
                "expires_in_minutes": SAS_URL_EXPIRY_MINUTES,
            }),
            status_code=200,
            mimetype="application/json",
        )

    except Exception as e:
        logging.error(f"generate_upload_url error: {str(e)}")
        return func.HttpResponse(
            json.dumps({"error": str(e)}),
            status_code=500,
            mimetype="application/json",
        )


# ==============================================================================
# BLOB TRIGGER - Automatic File Processing
# ==============================================================================
@app.blob_trigger(
    arg_name="blob",
    path="uploads/{name}",
    connection="AzureWebJobsStorage",
    data_type="binary",
)
def process_blob(blob: func.InputStream):
    """
    Automatic processing triggered by file uploads to Azure Blob Storage.
    
    This function is automatically triggered when files are uploaded to the
    'uploads' container. It processes both text documents and Excel files
    through their respective LangGraph pipelines and saves results to the
    'results' container.
    
    Supported File Types:
        - Text: .txt, .docx
        - Excel: .xlsx, .xls
        
    Output:
        - Enhanced .docx files (for text documents)
        - Processing metadata JSON files (for all files)
        - Error logs (if processing fails)
        
    Args:
        blob: Input stream of the uploaded file
    """
    blob_name = blob.name
    filename = blob_name.split("/")[-1]

    logging.info(f"Blob trigger activated for: {blob_name}")
    
    account_name = os.environ["STORAGE_ACCOUNT_NAME"]
    account_key = os.environ["STORAGE_ACCOUNT_KEY"]
    base_name, _ = os.path.splitext(filename)
    result_blob_name = f"{base_name}.metadata.json"

    try:
        # Validate supported file types
        if not filename.endswith(ALLOWED_ALL_EXTENSIONS):
            logging.warning(f"Unsupported file type: {filename} - skipping")
            return

        raw_bytes = blob.read()
        
        # Route to appropriate pipeline based on file extension
        if filename.lower().endswith(ALLOWED_EXCEL_EXTENSIONS):
            logging.info(f"Processing as Excel file: {filename}")
            result = run_excel_langgraph_pipeline(raw_bytes, filename)
            processing_result = {
                "success": result.get("success", False),
                "source_blob": blob_name,
                "metadata_blob": f"results/{result_blob_name}",
                "filename": filename,
                "title": result.get("title", ""),
                "summary": result.get("summary", ""),
                # Note: parsed_data excluded from metadata to reduce size
                "token_usage": result.get("token_usage", {}),
                "ragas_scores": result.get("ragas_scores", {}),
                "workflow_info": result.get("workflow_info", {}),
                "processing_timestamp": datetime.now(timezone.utc).isoformat(),
            }

        else:
            # Text/Word document processing
            text = _extract_text(raw_bytes, filename)
            logging.info(f"Extracted {len(text)} characters from: {filename}")

            if not text.strip():
                logging.warning(f"Empty document detected: {filename} - skipping")
                return

            logging.info("Executing text processing pipeline...")
            final_state = _run_agent_pipeline(text, filename)
            processing_result = {
                "success": True,
                "source_blob": blob_name,
                "result_blob": f"results/{base_name}_improved.docx",
                "metadata_blob": f"results/{result_blob_name}",
                "original_text": final_state.get("text", ""),
                "enhanced_text": final_state.get("enhanced_text", ""),
                "title": final_state.get("title", ""),
                "summary": final_state.get("summary", ""),
                "synonyms_found": len(final_state.get("synonyms", {})),
                "synonyms_applied": final_state.get("synonyms", {}),
                "token_usage": final_state.get("token_usage", {}),
                "parsed_data": final_state.get("parsed_data", {}),
                "workflow_info": {
                    "agents_used": [
                        "parse_text",
                        "find_equivalents",
                        "generate_summary",
                        "consolidate_text"
                    ],
                    "workflow_type": "langgraph_chat_completions_api_pipeline",
                    "version": WORKFLOW_VERSION,
                    "langgraph_enabled": True,
                    "azure_openai_chat_completions_enabled": True,
                    "processing_timestamp": datetime.now(timezone.utc).isoformat()
                },
            }

        # Initialize blob storage client
        blob_service = BlobServiceClient(
            account_url=f"https://{account_name}.blob.core.windows.net",
            credential=account_key,
        )
        results_container = blob_service.get_container_client(STORAGE_CONTAINER_RESULTS)

        # Create and upload enhanced .docx for text documents
        if not filename.lower().endswith(ALLOWED_EXCEL_EXTENSIONS):
            improved_docx_name = f"{base_name}_improved.docx"
            enhanced_text = (
                final_state.get("enhanced_text", "") or 
                final_state.get("text", "")
            )
            improved_docx_bytes = _build_docx_bytes_from_text(enhanced_text)
            results_container.upload_blob(
                name=improved_docx_name,
                data=improved_docx_bytes,
                overwrite=True,
                content_settings=ContentSettings(
                    content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                ),
            )
            logging.info(f"Enhanced document saved: results/{improved_docx_name}")

        # Upload metadata JSON for all file types
        results_container.upload_blob(
            name=result_blob_name,
            data=json.dumps(processing_result, ensure_ascii=False, indent=2),
            overwrite=True,
            content_settings=ContentSettings(content_type="application/json"),
        )
        logging.info(f"Metadata saved: results/{result_blob_name}")
        logging.info(f"Processing completed successfully: {filename}")

    except Exception as e:
        # Comprehensive error handling and logging
        error_payload = {
            "success": False,
            "source_blob": blob_name,
            "error": str(e),
            "traceback": traceback.format_exc(),
            "workflow_info": {
                "workflow_type": "langgraph_completions_api_pipeline",
                "version": WORKFLOW_VERSION,
                "error_timestamp": datetime.now(timezone.utc).isoformat()
            },
        }

        logging.exception(f"Processing failed for: {blob_name}")

        # Persist error details for debugging
        try:
            blob_service = BlobServiceClient(
                account_url=f"https://{account_name}.blob.core.windows.net",
                credential=account_key,
            )
            results_container = blob_service.get_container_client(STORAGE_CONTAINER_RESULTS)
            results_container.upload_blob(
                name=result_blob_name,
                data=json.dumps(error_payload, ensure_ascii=False, indent=2),
                overwrite=True,
            )
            logging.info(f"Error details saved to: results/{result_blob_name}")
        except Exception:
            logging.exception(f"Failed to save error details for: {blob_name}")
