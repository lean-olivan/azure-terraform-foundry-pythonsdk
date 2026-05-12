"""
Excel Parser Agent
Parses Excel files and extracts structured data and text content for analysis.
"""

import logging
import re
import openpyxl
from io import BytesIO
from typing import TypedDict, Dict, Any, List, Optional


# ============================================================================
# State Definition
# ============================================================================

class ExcelAgentState(TypedDict):
    """State structure for Excel processing pipeline"""
    text: str
    parsed_data: dict
    synonyms: dict
    title: str
    summary: str
    enhanced_text: str
    pdf_content: str
    token_usage: dict
    excel_content: bytes
    filename: str


# ============================================================================
# Constants
# ============================================================================

# Header detection keywords
HEADER_KEYWORDS = [
    'date', 'name', 'id', 'amount', 'price', 
    'status', 'type', 'category', 'description'
]

# Date format patterns
DATE_PATTERNS = [
    r'\d{4}-\d{2}-\d{2}',      # YYYY-MM-DD
    r'\d{2}/\d{2}/\d{4}',      # MM/DD/YYYY
    r'\d{2}-\d{2}-\d{4}',      # MM-DD-YYYY
]

# Complexity thresholds
COMPLEXITY_SIMPLE_SHEETS = 2
COMPLEXITY_SIMPLE_CELLS = 100
COMPLEXITY_MEDIUM_SHEETS = 5
COMPLEXITY_MEDIUM_CELLS = 500


# ============================================================================
# Main Parser Function
# ============================================================================

def parse_excel_agent(state: ExcelAgentState) -> ExcelAgentState:
    """
    Parse and analyze Excel file content.
    
    Extracts structured data, metadata, and text content from Excel files
    for downstream processing.
    
    Args:
        state: Current pipeline state with excel_content and filename
        
    Returns:
        Updated state with parsed_data and extracted text
    """
    excel_content = state.get("excel_content", b"")
    filename = state.get("filename", "unknown.xlsx")
    
    # Validate input
    if not excel_content:
        logging.error("No Excel content provided to parse_excel_agent")
        return _set_error_state(state, "No Excel content provided", filename)
    
    try:
        # Parse the Excel file
        parsed_result = _parse_excel_file(excel_content, filename)
        
        # Update state with results
        state["parsed_data"] = parsed_result["parsed_data"]
        state["text"] = parsed_result["extracted_text"]
        
        logging.info(
            f"Successfully parsed Excel file '{filename}': "
            f"{parsed_result['parsed_data']['total_sheets']} sheets, "
            f"{parsed_result['parsed_data']['total_rows']} rows, "
            f"{parsed_result['parsed_data']['total_cells']} cells"
        )
        
    except Exception as e:
        logging.error(f"Error parsing Excel file '{filename}': {str(e)}")
        return _set_error_state(state, str(e), filename)
    
    return state


# ============================================================================
# Excel Parsing Functions
# ============================================================================

def _parse_excel_file(excel_content: bytes, filename: str) -> Dict[str, Any]:
    """
    Parse Excel file and extract all data and metadata.
    
    Args:
        excel_content: Raw bytes of Excel file
        filename: Name of the file
        
    Returns:
        Dictionary with parsed_data and extracted_text
    """
    excel_file = BytesIO(excel_content)
    
    # Get sheet names using openpyxl
    sheet_names = _get_sheet_names(excel_file)
    
    # Parse each sheet
    sheets_data = {}
    text_contents = []
    total_stats = {"rows": 0, "columns": 0, "cells": 0}
    
    for sheet_name in sheet_names:
        sheet_result = _parse_sheet(excel_file, sheet_name)
        sheets_data[sheet_name] = sheet_result["data"]
        text_contents.append(sheet_result["text"])
        
        # Accumulate statistics
        total_stats["rows"] += sheet_result["data"]["rows"]
        total_stats["columns"] += sheet_result["data"]["columns"]
        total_stats["cells"] += sheet_result["data"]["non_empty_cells"]
    
    # Combine all text
    combined_text = "\n\n".join(text_contents)
    
    # Create comprehensive parsed data structure
    parsed_data = {
        "file_type": "excel",
        "filename": filename,
        "total_sheets": len(sheet_names),
        "sheet_names": sheet_names,
        "total_rows": total_stats["rows"],
        "total_columns": total_stats["columns"],
        "total_cells": total_stats["cells"],
        "sheets": sheets_data,
        "extracted_text": combined_text,
        "text_length": len(combined_text),
        "file_analysis": _analyze_file_structure(sheets_data)
    }
    
    return {
        "parsed_data": parsed_data,
        "extracted_text": combined_text
    }


def _get_sheet_names(excel_file: BytesIO) -> List[str]:
    """
    Extract sheet names from Excel file.
    
    Args:
        excel_file: BytesIO object containing Excel data
        
    Returns:
        List of sheet names
    """
    excel_file.seek(0)  # Reset to beginning
    workbook = openpyxl.load_workbook(excel_file, read_only=True)
    sheet_names = workbook.sheetnames
    workbook.close()
    return sheet_names


def _parse_sheet(excel_file: BytesIO, sheet_name: str) -> Dict[str, Any]:
    """
    Parse a single Excel sheet using openpyxl only.
    
    Args:
        excel_file: BytesIO object containing Excel data
        sheet_name: Name of sheet to parse
        
    Returns:
        Dictionary with sheet data and extracted text
    """
    try:
        excel_file.seek(0)  # Reset to beginning
        
        # Load workbook with openpyxl (data_only=True to get computed values)
        wb = openpyxl.load_workbook(excel_file, data_only=True)
        ws = wb[sheet_name]
        
        # Extract all data from the sheet
        data = []
        for row in ws.iter_rows(values_only=True):
            data.append(list(row))
        
        wb.close()
        
        # Check if sheet is empty
        if not data or all(all(cell is None for cell in row) for row in data):
            logging.warning(f"Sheet '{sheet_name}' has no data")
            return {
                "data": {
                    "sheet_name": sheet_name,
                    "rows": 0,
                    "columns": 0,
                    "non_empty_cells": 0,
                    "data": [],
                    "has_headers": False,
                    "data_types": {},
                },
                "text": f"--- Sheet: {sheet_name} (Empty) ---"
            }
        
        # Calculate statistics
        sheet_rows = len(data)
        sheet_columns = max(len(row) for row in data) if data else 0
        
        # Normalize row lengths (ensure all rows have same number of columns)
        normalized_data = []
        for row in data:
            normalized_row = list(row) + [None] * (sheet_columns - len(row))
            normalized_data.append(normalized_row)
        
        # Count non-empty cells
        non_empty_cells = sum(
            1 for row in normalized_data 
            for cell in row 
            if cell is not None and str(cell).strip()
        )
        
        # Convert data to strings for storage
        string_data = [
            [str(cell) if cell is not None else "" for cell in row]
            for row in normalized_data
        ]
        
        # Detect headers and column types
        has_headers = _detect_headers_openpyxl(normalized_data)
        data_types = _detect_column_types_openpyxl(normalized_data)
        
        # Create structured sheet data
        sheet_data = {
            "sheet_name": sheet_name,
            "rows": sheet_rows,
            "columns": sheet_columns,
            "non_empty_cells": non_empty_cells,
            "data": string_data,
            "has_headers": has_headers,
            "data_types": data_types,
        }
        
        # Extract text content
        text_content = _extract_sheet_text_openpyxl(normalized_data, sheet_name)
        
        logging.info(
            f"Parsed sheet '{sheet_name}': "
            f"{sheet_rows}x{sheet_columns}, {non_empty_cells} non-empty cells"
        )
        
        return {"data": sheet_data, "text": text_content}
        
    except Exception as e:
        logging.error(f"Error parsing sheet '{sheet_name}': {str(e)}")
        return {
            "data": {
                "sheet_name": sheet_name,
                "error": str(e),
                "rows": 0,
                "columns": 0,
                "non_empty_cells": 0,
                "data": []
            },
            "text": f"--- Sheet: {sheet_name} (Error: {str(e)}) ---"
        }


# ============================================================================
# Data Type Detection Functions (openpyxl-based)
# ============================================================================

def _detect_headers_openpyxl(data: List[List[Any]]) -> bool:
    """
    Detect if the first row likely contains column headers.
    
    Args:
        data: 2D list of cell values from openpyxl
        
    Returns:
        True if first row appears to be headers
    """
    if len(data) < 2:
        return False
    
    first_row = data[0]
    
    # Check for common header keywords
    for cell_value in first_row:
        if cell_value is not None:
            cell_str = str(cell_value).lower()
            if any(keyword in cell_str for keyword in HEADER_KEYWORDS):
                return True
    
    # Check if first row has different pattern than subsequent rows
    # Compare if first row is all strings while other rows have mixed types
    try:
        first_row_types = set(type(cell) for cell in first_row if cell is not None)
        
        if len(data) > 1:
            # Get types from a sample of other rows
            other_rows_types = set()
            for row in data[1:min(6, len(data))]:
                other_rows_types.update(type(cell) for cell in row if cell is not None)
            
            # If first row is all strings but other rows have numbers, likely headers
            if len(first_row_types) == 1 and str in first_row_types and len(other_rows_types) > 1:
                return True
    except Exception as e:
        logging.debug(f"Error detecting headers: {e}")
    
    return False


def _detect_column_types_openpyxl(data: List[List[Any]]) -> Dict[str, str]:
    """
    Detect data types for each column.
    
    Args:
        data: 2D list of cell values from openpyxl
        
    Returns:
        Dictionary mapping column indices to detected types
    """
    column_types = {}
    
    if not data:
        return column_types
    
    num_columns = len(data[0]) if data else 0
    
    for col_idx in range(num_columns):
        # Extract column values (skip None/empty)
        column_values = []
        for row in data:
            if col_idx < len(row) and row[col_idx] is not None:
                value_str = str(row[col_idx]).strip()
                if value_str:
                    column_values.append(value_str)
        
        if not column_values:
            column_types[f"column_{col_idx}"] = "empty"
            continue
        
        # Sample first few non-empty values (up to 5)
        sample_values = column_values[:5]
        
        # Detect type based on patterns
        detected_type = _classify_data_type(sample_values)
        column_types[f"column_{col_idx}"] = detected_type
    
    return column_types


def _classify_data_type(sample_values: List[str]) -> str:
    """
    Classify data type based on sample values.
    
    Args:
        sample_values: List of sample values to analyze
        
    Returns:
        Detected data type as string
    """
    non_empty_values = [v for v in sample_values if v.strip()]
    
    if not non_empty_values:
        return "empty"
    
    # Check each type in order of specificity
    if all(_is_date(val) for val in non_empty_values):
        return "date"
    elif all(_is_currency(val) for val in non_empty_values):
        return "currency"
    elif all(_is_percentage(val) for val in non_empty_values):
        return "percentage"
    elif all(_is_numeric(val) for val in non_empty_values):
        return "numeric"
    else:
        return "text"


def _is_date(value: str) -> bool:
    """Check if value matches date patterns"""
    return any(re.match(pattern, value.strip()) for pattern in DATE_PATTERNS)


def _is_numeric(value: str) -> bool:
    """Check if value is numeric"""
    try:
        float(value.replace(',', '').strip())
        return True
    except ValueError:
        return False


def _is_currency(value: str) -> bool:
    """Check if value looks like currency"""
    value = value.strip()
    return (value.startswith(('$', '€', '£', '¥')) or 
            value.endswith(('USD', 'EUR', 'GBP', 'JPY')))


def _is_percentage(value: str) -> bool:
    """Check if value looks like a percentage"""
    return value.strip().endswith('%')


# ============================================================================
# Text Extraction Functions (openpyxl-based)
# ============================================================================

def _extract_sheet_text_openpyxl(data: List[List[Any]], sheet_name: str) -> str:
    """
    Extract readable text content from a sheet.
    
    Args:
        data: 2D list of cell values from openpyxl
        sheet_name: Name of the sheet
        
    Returns:
        Formatted text content
    """
    text_lines = [f"--- Sheet: {sheet_name} ---"]
    
    for row in data:
        # Join non-empty cells with pipe separator
        row_text = " | ".join([
            str(cell) for cell in row 
            if cell is not None and str(cell).strip()
        ])
        
        if row_text:
            text_lines.append(row_text)
    
    return "\n".join(text_lines)


# ============================================================================
# File Analysis Functions
# ============================================================================

def _analyze_file_structure(sheets_data: Dict[str, Any]) -> Dict[str, Any]:
    """
    Analyze the overall structure and content of the Excel file.
    
    Args:
        sheets_data: Dictionary of all parsed sheet data
        
    Returns:
        Dictionary with file analysis metadata
    """
    return {
        "complexity": _assess_complexity(sheets_data),
        "data_categories": _identify_data_categories(sheets_data),
        "structure_type": _identify_structure_type(sheets_data),
        "content_summary": _generate_content_summary(sheets_data)
    }


def _assess_complexity(sheets_data: Dict[str, Any]) -> str:
    """
    Assess the complexity level of the Excel file.
    
    Args:
        sheets_data: Dictionary of sheet data
        
    Returns:
        Complexity level: "Simple", "Medium", or "Complex"
    """
    total_sheets = len(sheets_data)
    total_cells = sum(sheet.get("non_empty_cells", 0) for sheet in sheets_data.values())
    
    if total_sheets <= COMPLEXITY_SIMPLE_SHEETS and total_cells <= COMPLEXITY_SIMPLE_CELLS:
        return "Simple"
    elif total_sheets <= COMPLEXITY_MEDIUM_SHEETS and total_cells <= COMPLEXITY_MEDIUM_CELLS:
        return "Medium"
    else:
        return "Complex"


def _identify_data_categories(sheets_data: Dict[str, Any]) -> List[str]:
    """
    Identify potential data categories based on content patterns.
    
    Args:
        sheets_data: Dictionary of sheet data
        
    Returns:
        List of identified categories
    """
    categories = []
    sheet_names_combined = " ".join([
        sheet.get("sheet_name", "").lower() 
        for sheet in sheets_data.values()
    ])
    
    # Check for category keywords
    category_keywords = {
        "Financial": ["financial", "budget", "revenue", "cost", "expense"],
        "Risk Management": ["risk", "assessment", "mitigation", "threat"],
        "Project Management": ["project", "task", "timeline", "milestone"],
        "Customer Data": ["customer", "client", "sales", "contact"],
        "Inventory": ["inventory", "stock", "product", "warehouse"],
        "HR/Personnel": ["employee", "staff", "payroll", "personnel"],
        "Compliance": ["compliance", "audit", "regulation", "policy"]
    }
    
    for category, keywords in category_keywords.items():
        if any(keyword in sheet_names_combined for keyword in keywords):
            categories.append(category)
    
    return categories if categories else ["General"]


def _identify_structure_type(sheets_data: Dict[str, Any]) -> str:
    """
    Identify the structural type of the Excel file.
    
    Args:
        sheets_data: Dictionary of sheet data
        
    Returns:
        Structure type description
    """
    sheet_count = len(sheets_data)
    
    if sheet_count == 1:
        return "Single Sheet"
    elif sheet_count <= 3:
        return "Multi-Sheet"
    else:
        return "Complex Workbook"


def _generate_content_summary(sheets_data: Dict[str, Any]) -> str:
    """
    Generate a brief text summary of the file content.
    
    Args:
        sheets_data: Dictionary of sheet data
        
    Returns:
        Summary text
    """
    sheet_names = [
        sheet.get("sheet_name", "Unknown") 
        for sheet in sheets_data.values()
    ]
    total_cells = sum(
        sheet.get("non_empty_cells", 0) 
        for sheet in sheets_data.values()
    )
    
    return (
        f"Excel workbook with {len(sheet_names)} sheet(s) "
        f"({', '.join(sheet_names)}) containing {total_cells} data cells"
    )


# ============================================================================
# Error Handling Functions
# ============================================================================

def _set_error_state(state: ExcelAgentState, error_message: str, filename: str) -> ExcelAgentState:
    """
    Set error information in state when parsing fails.
    
    Args:
        state: Current pipeline state
        error_message: Error description
        filename: Name of file that failed
        
    Returns:
        Updated state with error information
    """
    state["parsed_data"] = {
        "error": error_message,
        "file_type": "excel",
        "filename": filename,
        "total_sheets": 0,
        "sheet_names": [],
        "total_rows": 0,
        "total_columns": 0,
        "total_cells": 0,
        "sheets": {},
        "extracted_text": ""
    }
    state["text"] = ""
    return state
